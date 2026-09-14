"""
Output formats for converted Markdown: .md, .txt, .json, .html, .docx, .pdf.

`render(fmt, markdown, source_name=..., sidecar=...)` returns bytes ready for a
download button; `zip_bundle()` packs several renders into one archive.

Dependencies: `markdown` (HTML), `python-docx` (DOCX), `fpdf2` (PDF).
"""

from __future__ import annotations

import datetime as _dt
import html as _html
import io
import json
import os
import re
import zipfile
from html.parser import HTMLParser

# fmt key -> (label, mime, extension)
FORMATS: dict[str, tuple[str, str, str]] = {
    "md": ("Markdown (.md)", "text/markdown", ".md"),
    "txt": ("Plain text (.txt)", "text/plain", ".txt"),
    "json": ("JSON (.json)", "application/json", ".json"),
    "html": ("HTML (.html)", "text/html", ".html"),
    "docx": ("Word (.docx)", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
    "pdf": ("PDF (.pdf)", "application/pdf", ".pdf"),
}
DEFAULT_FORMATS = ["md"]

_MD_EXTENSIONS = ["tables", "fenced_code", "sane_lists", "toc", "attr_list"]


def _safe_stem(source_name: str) -> str:
    stem = os.path.splitext(os.path.basename(source_name or ""))[0]
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", stem).strip(" .")
    return stem or "document"


def filename_for(source_name: str, fmt: str) -> str:
    return _safe_stem(source_name) + FORMATS[fmt][2]


def title_for(markdown: str, source_name: str) -> str:
    m = re.search(r"^#\s+(.+?)\s*$", markdown or "", re.M)
    return m.group(1).strip() if m else _safe_stem(source_name)


# --------------------------------------------------------------------------- #
# HTML
# --------------------------------------------------------------------------- #
def to_html_fragment(markdown_text: str) -> str:
    import markdown

    return markdown.markdown(markdown_text or "", extensions=_MD_EXTENSIONS, output_format="html")


_HTML_CSS = """
body{font-family:Segoe UI,Helvetica,Arial,sans-serif;max-width:900px;margin:2rem auto;padding:0 1rem;line-height:1.55;color:#222}
h1,h2,h3{line-height:1.25} pre{background:#f5f5f5;padding:.8rem;overflow:auto;border-radius:6px}
code{background:#f5f5f5;padding:.1rem .3rem;border-radius:4px;font-family:Consolas,Menlo,monospace}
table{border-collapse:collapse;margin:1rem 0} th,td{border:1px solid #bbb;padding:.35rem .6rem;vertical-align:top}
th{background:#eee} blockquote{border-left:4px solid #ccc;margin:1rem 0;padding:.2rem 1rem;color:#555} hr{border:0;border-top:1px solid #ddd}
"""


def to_html(markdown_text: str, source_name: str) -> str:
    title = _html.escape(title_for(markdown_text, source_name))
    return (
        "<!DOCTYPE html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n"
        f"<title>{title}</title>\n<style>{_HTML_CSS}</style>\n</head>\n<body>\n"
        f"{to_html_fragment(markdown_text)}\n</body>\n</html>\n"
    )


# --------------------------------------------------------------------------- #
# Plain text
# --------------------------------------------------------------------------- #
class _TextExtractor(HTMLParser):
    _BLOCK = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr", "pre", "blockquote", "table", "ul", "ol", "hr", "br"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._in_cell = False

    def handle_starttag(self, tag, attrs):
        if tag in ("td", "th"):
            self._in_cell = True
            if self.parts and not self.parts[-1].endswith(("\n", "\t")):
                self.parts.append("\t")
        elif tag == "li":
            self.parts.append("\n- ")
        elif tag == "hr":
            self.parts.append("\n" + "-" * 40 + "\n")
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("td", "th"):
            self._in_cell = False
        elif tag in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "blockquote", "table", "tr"):
            self.parts.append("\n")

    def handle_data(self, data):
        self.parts.append(data if self._in_cell or "\n" in data else data)

    def text(self) -> str:
        text = "".join(self.parts)
        text = re.sub(r"[ \t]+\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip() + "\n"


def to_text(markdown_text: str) -> str:
    parser = _TextExtractor()
    parser.feed(to_html_fragment(markdown_text))
    parser.close()
    return parser.text()


# --------------------------------------------------------------------------- #
# JSON
# --------------------------------------------------------------------------- #
def to_json(markdown_text: str, source_name: str, sidecar: dict | None = None, extra: dict | None = None) -> str:
    payload = {
        "source_file": source_name,
        "generated_at": _dt.datetime.now().isoformat(timespec="seconds"),
        "title": title_for(markdown_text, source_name),
        "markdown": markdown_text,
        "plain_text": to_text(markdown_text),
        "sidecar": sidecar,
    }
    if extra:
        payload.update(extra)
    return json.dumps(payload, ensure_ascii=False, indent=2)


# --------------------------------------------------------------------------- #
# Block-level Markdown parser (shared by DOCX)
# --------------------------------------------------------------------------- #
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_LIST_RE = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_HR_RE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_INLINE_RE = re.compile(
    r"(\*\*[^*\n]+\*\*|__[^_\n]+__|`[^`\n]+`|!\[[^\]]*\]\([^)]*\)|\[[^\]\n]+\]\([^)]*\)|\*[^*\n]+\*|_[^_\n]+_)"
)


class _TableHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self._row = None
        self._cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append("\n")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            self._row.append("".join(self._cell).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._row:
                self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _split_md_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip().replace("\\|", "|") for c in s.split("|")]


def parse_blocks(markdown_text: str) -> list[tuple]:
    """
    Yield block tuples:
      ("heading", level, text) ("para", text) ("code", lang, text) ("quote", text)
      ("list", [(level, ordered, text), ...]) ("table", rows) ("hr",)
    """
    lines = (markdown_text or "").replace("\r\n", "\n").split("\n")
    blocks: list[tuple] = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```") or s.startswith("~~~"):
            fence = s[:3]
            lang = s[3:].strip()
            i += 1
            code: list[str] = []
            while i < n and not lines[i].strip().startswith(fence):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", lang, "\n".join(code)))
            continue
        if s.lower().startswith("<table"):
            buf: list[str] = []
            while i < n:
                buf.append(lines[i])
                if "</table>" in lines[i].lower():
                    i += 1
                    break
                i += 1
            parser = _TableHTMLParser()
            parser.feed("\n".join(buf))
            parser.close()
            blocks.append(("table", parser.rows))
            continue
        m = _HEADING_RE.match(s)
        if m:
            blocks.append(("heading", len(m.group(1)), m.group(2)))
            i += 1
            continue
        if _HR_RE.match(s):
            blocks.append(("hr",))
            i += 1
            continue
        if s.startswith("|") and i + 1 < n and _TABLE_SEP_RE.match(lines[i + 1].strip()):
            rows = [_split_md_row(lines[i])]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_split_md_row(lines[i]))
                i += 1
            blocks.append(("table", rows))
            continue
        if s.startswith(">"):
            quote: list[str] = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(("quote", "\n".join(quote)))
            continue
        m = _LIST_RE.match(line)
        if m:
            items: list[tuple] = []
            while i < n:
                m = _LIST_RE.match(lines[i])
                if not m:
                    break
                indent = len(m.group(1).replace("\t", "    "))
                items.append((min(indent // 2, 2), m.group(2)[0].isdigit(), m.group(3).strip()))
                i += 1
            blocks.append(("list", items))
            continue
        para: list[str] = [s]
        i += 1
        while i < n:
            nxt = lines[i]
            ns = nxt.strip()
            if (
                not ns
                or ns.startswith(("```", "~~~", "|", ">", "#", "<table"))
                or _LIST_RE.match(nxt)
                or _HR_RE.match(ns)
            ):
                break
            para.append(ns)
            i += 1
        blocks.append(("para", " ".join(para)))
    return blocks


def _inline_runs(text: str) -> list[tuple[str, dict]]:
    """Split inline Markdown into (text, {bold, italic, code}) runs."""
    runs: list[tuple[str, dict]] = []
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos : m.start()], {}))
        tok = m.group(0)
        if tok.startswith("**") or tok.startswith("__"):
            runs.append((tok[2:-2], {"bold": True}))
        elif tok.startswith("`"):
            runs.append((tok[1:-1], {"code": True}))
        elif tok.startswith("!["):
            alt = re.match(r"!\[([^\]]*)\]", tok).group(1)
            runs.append((f"[image: {alt or 'image'}]", {"italic": True}))
        elif tok.startswith("["):
            mm = re.match(r"\[([^\]]+)\]\(([^)]*)\)", tok)
            label, url = mm.group(1), mm.group(2)
            runs.append((label, {}))
            if url and url != label:
                runs.append((f" ({url})", {"italic": True}))
        else:
            runs.append((tok[1:-1], {"italic": True}))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], {}))
    return runs


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def to_docx(markdown_text: str, source_name: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = title_for(markdown_text, source_name)

    def add_runs(paragraph, text: str):
        for chunk, style in _inline_runs(text):
            run = paragraph.add_run(chunk)
            if style.get("bold"):
                run.bold = True
            if style.get("italic"):
                run.italic = True
            if style.get("code"):
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)

    for block in parse_blocks(markdown_text):
        kind = block[0]
        if kind == "heading":
            doc.add_heading(re.sub(r"[*_`]", "", block[2]), level=min(block[1], 9))
        elif kind == "para":
            add_runs(doc.add_paragraph(), block[1])
        elif kind == "quote":
            add_runs(doc.add_paragraph(style="Quote"), block[1])
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(block[2])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif kind == "hr":
            doc.add_paragraph("─" * 30)
        elif kind == "list":
            for level, ordered, text in block[1]:
                base = "List Number" if ordered else "List Bullet"
                style = base if level == 0 else f"{base} {min(level + 1, 3)}"
                try:
                    p = doc.add_paragraph(style=style)
                except KeyError:
                    p = doc.add_paragraph(style=base)
                add_runs(p, text)
        elif kind == "table":
            rows = block[1]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
_WIN_FONTS = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")


def _find_font_set() -> dict | None:
    """Locate a Unicode TTF family (regular/bold/italic/bold-italic), a mono font and an emoji font."""
    candidates = [
        {"": "segoeui.ttf", "B": "segoeuib.ttf", "I": "segoeuii.ttf", "BI": "segoeuiz.ttf"},
        {"": "arial.ttf", "B": "arialbd.ttf", "I": "ariali.ttf", "BI": "arialbi.ttf"},
        {"": "DejaVuSans.ttf", "B": "DejaVuSans-Bold.ttf", "I": "DejaVuSans-Oblique.ttf", "BI": "DejaVuSans-BoldOblique.ttf"},
    ]
    search_dirs = [_WIN_FONTS, "/usr/share/fonts/truetype/dejavu", "/Library/Fonts", "/System/Library/Fonts"]
    try:
        import matplotlib

        search_dirs.append(os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf"))
    except Exception:
        pass

    def locate(filename: str) -> str | None:
        for d in search_dirs:
            p = os.path.join(d, filename)
            if os.path.exists(p):
                return p
        return None

    for fam in candidates:
        regular = locate(fam[""])
        if not regular:
            continue
        found = {"": regular}
        for style in ("B", "I", "BI"):
            p = locate(fam[style])
            if p:
                found[style] = p
        mono = locate("consola.ttf") or locate("cour.ttf") or locate("DejaVuSansMono.ttf")
        emoji = locate("seguiemj.ttf")
        return {"body": found, "mono": mono, "emoji": emoji}
    return None


def _latin1_safe(text: str) -> str:
    return text.encode("latin-1", "replace").decode("latin-1")


def to_pdf(markdown_text: str, source_name: str) -> bytes:
    from fpdf import FPDF

    fragment = to_html_fragment(markdown_text)
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_title(title_for(markdown_text, source_name))
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(15, 15, 15)
    pdf.add_page()

    fonts = _find_font_set()
    body_family, mono_family = "helvetica", "courier"
    if fonts:
        for style, path in fonts["body"].items():
            pdf.add_font("Body", style=style, fname=path)
        body_family = "Body"
        if fonts["mono"]:
            pdf.add_font("Mono", style="", fname=fonts["mono"])
            mono_family = "Mono"
        if fonts["emoji"]:
            pdf.add_font("Emoji", style="", fname=fonts["emoji"])
            pdf.set_fallback_fonts(["Emoji"])
        else:
            # No emoji font: drop astral-plane glyphs rather than printing boxes.
            fragment = re.sub(r"[\U00010000-\U0010FFFF]", "", fragment)
    else:
        fragment = _latin1_safe(fragment)

    pdf.set_font(body_family, size=11)
    try:
        pdf.write_html(fragment, font_family=body_family, pre_code_font=mono_family)
    except Exception:
        # Last resort: plain text so the user still gets a file.
        pdf.set_font(body_family, size=11)
        text = to_text(markdown_text)
        if not fonts:
            text = _latin1_safe(text)
        pdf.multi_cell(0, 6, text)
    return bytes(pdf.output())


# --------------------------------------------------------------------------- #
# Dispatcher + bundling
# --------------------------------------------------------------------------- #
def render(fmt: str, markdown_text: str, *, source_name: str, sidecar: dict | None = None, extra: dict | None = None) -> bytes:
    """Render Markdown into the requested format and return bytes."""
    if fmt == "md":
        return (markdown_text or "").encode("utf-8")
    if fmt == "txt":
        return to_text(markdown_text).encode("utf-8")
    if fmt == "json":
        return to_json(markdown_text, source_name, sidecar, extra).encode("utf-8")
    if fmt == "html":
        return to_html(markdown_text, source_name).encode("utf-8")
    if fmt == "docx":
        return to_docx(markdown_text, source_name)
    if fmt == "pdf":
        return to_pdf(markdown_text, source_name)
    raise ValueError(f"Unknown format {fmt!r}; expected one of {list(FORMATS)}")


def zip_bundle(items: list[tuple[str, bytes]]) -> bytes:
    """Pack (filename, bytes) pairs into a zip archive; duplicate names get a numeric suffix."""
    buf = io.BytesIO()
    seen: dict[str, int] = {}
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in items:
            count = seen.get(name, 0)
            seen[name] = count + 1
            if count:
                stem, ext = os.path.splitext(name)
                name = f"{stem} ({count}){ext}"
            zf.writestr(name, data)
    return buf.getvalue()
