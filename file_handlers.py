"""
Per-file-type conversion handlers for the MarkItDown app.

Every handler returns a `ConversionResult` (Markdown + optional structured
sidecar). The Streamlit app decides *which* handler to call and applies the
user's mode choices (Claude / Ollama / Unlimited-OCR); this module stays free
of Streamlit so it can be unit-tested directly.

Capabilities here:
  * classification of uploads by extension
  * subtitles (.srt/.vtt), presentations (.pptx), spreadsheets (.xlsx/.csv),
    e-mail (.eml), archives (.zip), extended images (.heic/.webp/.tiff, incl.
    multi-page TIFF), Tesseract OCR
  * MarkItDown default conversion
  * URL conversion (Wikipedia / YouTube / generic with 3 fallbacks)
  * AI enhancement of raw Markdown via Claude or Ollama
"""

from __future__ import annotations

import asyncio
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser

import requests

# --------------------------------------------------------------------------- #
# Extension registry
# --------------------------------------------------------------------------- #
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".jfif", ".webp", ".tiff", ".tif", ".heic", ".heif"}
MULTIPAGE_IMAGE_EXTS = {".tiff", ".tif"}
AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".ogg", ".flac"}
VIDEO_EXTS = {".mp4", ".mkv", ".avi", ".mov"}
SUBTITLE_EXTS = {".srt", ".vtt"}
PRESENTATION_EXTS = {".pptx"}
SPREADSHEET_EXTS = {".xlsx", ".csv"}
EMAIL_EXTS = {".eml"}
ARCHIVE_EXTS = {".zip"}
TEXT_PREVIEW_EXTS = {".txt", ".md", ".csv", ".srt", ".vtt", ".eml", ".html", ".htm", ".json", ".xml", ".log", ".py", ".yaml", ".yml"}

MAX_ARCHIVE_FILES = 50
MAX_ARCHIVE_DEPTH = 2
MAX_TABLE_ROWS = 2000
MAX_IMAGE_PAGES = 200

# Document-mode labels shared with the app (kept here so handlers can compare).
DOC_MODE_DEFAULT = "⚡ MarkItDown Default"
DOC_MODE_CLAUDE = "🤖 Anthropic Claude AI"
DOC_MODE_OLLAMA = "🦙 Llama (Local, Free)"


@dataclass
class ConversionResult:
    markdown: str
    source_name: str
    kind: str                       # document | image | audio | video | subtitle | presentation | spreadsheet | email | archive | url
    sidecar: dict | None = None     # structured data (tables, cues, segments, ...) for the JSON output
    notes: list[str] = field(default_factory=list)   # informational messages for the UI
    children: list["ConversionResult"] = field(default_factory=list)  # archive members


def ext_of(name: str) -> str:
    return os.path.splitext(name or "")[1].lower()


def classify(name: str) -> str:
    """Map a filename onto a handler kind (see ConversionResult.kind)."""
    ext = ext_of(name)
    if ext in IMAGE_EXTS:
        return "image"
    if ext in AUDIO_EXTS:
        return "audio"
    if ext in VIDEO_EXTS:
        return "video"
    if ext in SUBTITLE_EXTS:
        return "subtitle"
    if ext in PRESENTATION_EXTS:
        return "presentation"
    if ext in SPREADSHEET_EXTS:
        return "spreadsheet"
    if ext in EMAIL_EXTS:
        return "email"
    if ext in ARCHIVE_EXTS:
        return "archive"
    return "document"


# --------------------------------------------------------------------------- #
# Small shared helpers
# --------------------------------------------------------------------------- #
def _stem(name: str) -> str:
    return os.path.splitext(os.path.basename(name))[0]


def _md_escape_cell(value) -> str:
    text = "" if value is None else str(value)
    if text.lower() == "nan":
        text = ""
    return text.replace("|", "\\|").replace("\n", " ").strip()


def rows_to_markdown(rows: list[list], header: bool = True) -> str:
    """Render a list of rows as a Markdown pipe table."""
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [[_md_escape_cell(c) for c in list(r) + [""] * (width - len(r))] for r in rows]
    head = norm[0] if header else [f"Col {i + 1}" for i in range(width)]
    body = norm[1:] if header else norm
    lines = ["| " + " | ".join(head) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


class _HTMLTextExtractor(HTMLParser):
    """Very small HTML → text fallback (used only if MarkItDown is unavailable)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in ("p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)

    def text(self) -> str:
        return re.sub(r"\n{3,}", "\n\n", "".join(self.parts)).strip()


def html_to_markdown(html: str) -> str:
    """Convert an HTML string to Markdown with MarkItDown (fallback: plain text)."""
    tmp_path = None
    try:
        from markitdown import MarkItDown

        with tempfile.NamedTemporaryFile(delete=False, suffix=".html", mode="w", encoding="utf-8") as tmp:
            tmp.write(html)
            tmp_path = tmp.name
        return MarkItDown().convert(tmp_path).text_content
    except Exception:
        parser = _HTMLTextExtractor()
        parser.feed(html)
        return parser.text()
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.remove(tmp_path)


# --------------------------------------------------------------------------- #
# MarkItDown default
# --------------------------------------------------------------------------- #
def convert_with_markitdown(path: str, llm_client=None, llm_model: str | None = None) -> str:
    from markitdown import MarkItDown

    md = MarkItDown(llm_client=llm_client, llm_model=llm_model) if llm_client else MarkItDown()
    return md.convert(path).text_content


# --------------------------------------------------------------------------- #
# Subtitles (.srt / .vtt)
# --------------------------------------------------------------------------- #
_SUB_TAG_RE = re.compile(r"<[^>]+>|\{\\[^}]*\}")


def _sub_time_to_seconds(ts: str) -> float:
    ts = ts.strip().replace(",", ".")
    parts = ts.split(":")
    try:
        if len(parts) == 3:
            h, m, s = parts
            return int(h) * 3600 + int(m) * 60 + float(s)
        if len(parts) == 2:
            m, s = parts
            return int(m) * 60 + float(s)
        return float(ts)
    except ValueError:
        return 0.0


def _fmt_ts(seconds: float) -> str:
    seconds = max(0, int(round(seconds)))
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def parse_subtitles(text: str) -> list[dict]:
    """Parse SRT or WebVTT into cues: [{"index", "start", "end", "text"}, ...]."""
    text = text.lstrip("\ufeff").replace("\r\n", "\n").replace("\r", "\n")
    cues: list[dict] = []
    for block in re.split(r"\n\s*\n", text):
        lines = [ln for ln in block.strip().split("\n") if ln.strip()]
        if not lines:
            continue
        if lines[0].startswith(("WEBVTT", "NOTE", "STYLE", "REGION")):
            continue
        time_idx = next((i for i, ln in enumerate(lines) if "-->" in ln), None)
        if time_idx is None:
            continue
        start_s, end_s = lines[time_idx].split("-->", 1)
        end_s = end_s.strip().split(" ")[0]  # drop VTT cue settings (position:.. line:..)
        body = " ".join(_SUB_TAG_RE.sub("", ln).strip() for ln in lines[time_idx + 1 :]).strip()
        if not body:
            continue
        index = lines[0].strip() if time_idx == 1 else str(len(cues) + 1)
        cues.append(
            {
                "index": index,
                "start": round(_sub_time_to_seconds(start_s), 3),
                "end": round(_sub_time_to_seconds(end_s), 3),
                "text": body,
            }
        )
    return cues


def convert_subtitles(path: str, name: str) -> ConversionResult:
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        cues = parse_subtitles(fh.read())
    fmt = ext_of(name).lstrip(".").upper()
    lines = [f"# Subtitles: {name}", "", f"> Format: {fmt} · {len(cues)} cues", ""]
    if cues:
        lines += ["## Transcript", ""]
        # Merge cues into paragraphs at gaps > 2 s
        para: list[str] = []
        last_end = None
        for c in cues:
            if para and last_end is not None and c["start"] - last_end > 2.0:
                lines += [" ".join(para), ""]
                para = []
            para.append(c["text"])
            last_end = c["end"]
        if para:
            lines += [" ".join(para), ""]
        lines += ["## Timed cues", ""]
        lines += [f"- **[{_fmt_ts(c['start'])} → {_fmt_ts(c['end'])}]** {c['text']}" for c in cues]
    else:
        lines.append("_No cues found._")
    return ConversionResult("\n".join(lines).strip() + "\n", name, "subtitle", sidecar={"format": fmt, "cue_count": len(cues), "cues": cues})


# --------------------------------------------------------------------------- #
# Presentations (.pptx)
# --------------------------------------------------------------------------- #
def convert_pptx(path: str, name: str) -> ConversionResult:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    slides_meta: list[dict] = []
    lines = [f"# {_stem(name)}", "", f"> Source: {name} · {len(prs.slides)} slide(s)", ""]

    def walk(shapes, out: list[str], meta: dict, title_id):
        for shape in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
            if shape.shape_id == title_id:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, out, meta, title_id)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                meta["tables"].append(rows)
                out += ["", rows_to_markdown(rows), ""]
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                meta["images"] += 1
                out.append(f"*[image: {shape.name}]*")
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        out.append(f"{'  ' * para.level}- {text}")

    for i, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        meta = {"slide": i, "title": title, "tables": [], "images": 0, "notes": ""}
        lines.append(f"## Slide {i}" + (f": {title}" if title else ""))
        lines.append("")
        body: list[str] = []
        walk(slide.shapes, body, meta, title_shape.shape_id if title_shape is not None else None)
        lines += body
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                meta["notes"] = notes
                lines += ["", "> **Speaker notes:** " + notes.replace("\n", "\n> ")]
        lines.append("")
        slides_meta.append(meta)

    sidecar = {
        "slide_count": len(prs.slides),
        "slides": slides_meta,
        "tables": [{"slide": m["slide"], "index": j, "rows": t} for m in slides_meta for j, t in enumerate(m["tables"])],
    }
    return ConversionResult("\n".join(lines).strip() + "\n", name, "presentation", sidecar=sidecar)


# --------------------------------------------------------------------------- #
# Spreadsheets (.xlsx / .csv)
# --------------------------------------------------------------------------- #
def convert_spreadsheet(path: str, name: str, max_rows: int = MAX_TABLE_ROWS) -> ConversionResult:
    import pandas as pd

    ext = ext_of(name)
    if ext == ".csv":
        sheets = {"CSV": pd.read_csv(path, sep=None, engine="python", dtype=str, keep_default_na=False)}
    else:
        sheets = pd.read_excel(path, sheet_name=None, dtype=str)

    lines = [f"# {_stem(name)}", "", f"> Source: {name} · {len(sheets)} sheet(s)", ""]
    notes: list[str] = []
    sheets_meta: list[dict] = []
    for sheet_name, df in sheets.items():
        df = df.fillna("")
        rows = [list(map(str, df.columns))] + df.astype(str).values.tolist()
        lines.append(f"## {sheet_name}")
        lines.append("")
        lines.append(f"_{len(df)} rows × {len(df.columns)} columns_")
        lines.append("")
        shown = rows[: max_rows + 1]
        lines.append(rows_to_markdown(shown) if len(rows) > 1 else "_Empty sheet_")
        if len(rows) - 1 > max_rows:
            msg = f"Sheet '{sheet_name}': showing first {max_rows} of {len(rows) - 1} rows in Markdown (full data in JSON)."
            lines += ["", f"_{msg}_"]
            notes.append(msg)
        lines.append("")
        sheets_meta.append(
            {
                "name": str(sheet_name),
                "rows": len(df),
                "columns": len(df.columns),
                "header": rows[0],
                "data": rows[1:],
            }
        )
    sidecar = {"sheet_count": len(sheets), "sheets": sheets_meta, "tables": [{"sheet": s["name"], "rows": [s["header"]] + s["data"]} for s in sheets_meta]}
    return ConversionResult("\n".join(lines).strip() + "\n", name, "spreadsheet", sidecar=sidecar, notes=notes)


# --------------------------------------------------------------------------- #
# E-mail (.eml)
# --------------------------------------------------------------------------- #
def convert_email(path: str, name: str) -> ConversionResult:
    from email import policy
    from email.parser import BytesParser

    with open(path, "rb") as fh:
        msg = BytesParser(policy=policy.default).parse(fh)

    headers = {k: str(msg.get(k, "")) for k in ("From", "To", "Cc", "Bcc", "Date", "Subject", "Message-ID")}
    subject = headers["Subject"] or _stem(name)
    lines = [f"# {subject}", ""]
    lines.append(rows_to_markdown([["Header", "Value"]] + [[k, v] for k, v in headers.items() if v and k != "Message-ID"]))
    lines.append("")

    body_md = ""
    body = msg.get_body(preferencelist=("plain", "html"))
    if body is not None:
        content = body.get_content()
        body_md = html_to_markdown(content) if body.get_content_subtype() == "html" else content
    lines += ["## Body", "", body_md.strip() or "_(empty body)_", ""]

    attachments = []
    for part in msg.iter_attachments():
        payload = part.get_payload(decode=True) or b""
        attachments.append({"filename": part.get_filename() or "(unnamed)", "content_type": part.get_content_type(), "size_bytes": len(payload)})
    if attachments:
        lines += ["## Attachments", ""]
        lines += [f"- {a['filename']} ({a['content_type']}, {a['size_bytes']:,} bytes)" for a in attachments]
        lines.append("")

    sidecar = {"headers": headers, "attachments": attachments, "body_is_html": bool(body is not None and body.get_content_subtype() == "html")}
    return ConversionResult("\n".join(lines).strip() + "\n", name, "email", sidecar=sidecar)


# --------------------------------------------------------------------------- #
# Archives (.zip)
# --------------------------------------------------------------------------- #
def convert_archive(path: str, name: str, convert_fn, max_files: int = MAX_ARCHIVE_FILES, _depth: int = 0) -> ConversionResult:
    """
    Extract a .zip and convert each member with `convert_fn(member_path, member_name) -> ConversionResult`.
    Nested zips are handled up to MAX_ARCHIVE_DEPTH levels.
    """
    lines = [f"# Archive: {name}", ""]
    files_meta: list[dict] = []
    children: list[ConversionResult] = []
    notes: list[str] = []

    with tempfile.TemporaryDirectory(prefix="mdzip_") as work, zipfile.ZipFile(path) as zf:
        members = [
            m for m in zf.infolist()
            if not m.is_dir()
            and not m.filename.startswith("__MACOSX")
            and not os.path.basename(m.filename).startswith(".")
        ]
        lines.append(f"> {len(members)} file(s)" + (f" — first {max_files} converted" if len(members) > max_files else ""))
        lines.append("")
        root = os.path.realpath(work)
        for m in members[:max_files]:
            target = os.path.realpath(os.path.join(work, m.filename))
            if not target.startswith(root + os.sep):
                notes.append(f"Skipped unsafe path in archive: {m.filename}")
                continue
            zf.extract(m, work)
            lines.append(f"---\n\n## 📄 {m.filename}\n")
            entry = {"name": m.filename, "size_bytes": m.file_size, "status": "ok"}
            try:
                if ext_of(m.filename) in ARCHIVE_EXTS:
                    if _depth + 1 >= MAX_ARCHIVE_DEPTH:
                        raise RuntimeError(f"nested archives deeper than {MAX_ARCHIVE_DEPTH} levels are skipped")
                    res = convert_archive(target, m.filename, convert_fn, max_files, _depth + 1)
                else:
                    res = convert_fn(target, m.filename)
                children.append(res)
                lines.append(res.markdown.strip())
            except Exception as exc:
                entry["status"] = f"failed: {exc}"
                lines.append(f"⚠️ Conversion failed: {exc}")
            lines.append("")
            files_meta.append(entry)
        if len(members) > max_files:
            notes.append(f"Archive has {len(members)} files; only the first {max_files} were converted.")

    sidecar = {"file_count": len(files_meta), "files": files_meta, "members": [{"name": c.source_name, "kind": c.kind, "sidecar": c.sidecar} for c in children]}
    return ConversionResult("\n".join(lines).strip() + "\n", name, "archive", sidecar=sidecar, notes=notes, children=children)


# --------------------------------------------------------------------------- #
# Images (incl. .heic / .webp / multi-page .tiff)
# --------------------------------------------------------------------------- #
def _register_heif() -> None:
    try:
        from pillow_heif import register_heif_opener

        register_heif_opener()
    except ImportError:
        pass


def image_pages(path: str, out_dir: str, max_pages: int = MAX_IMAGE_PAGES) -> list[str]:
    """
    Normalise any supported image to one RGB PNG per page. Multi-page TIFFs
    yield one file per page; every other format yields a single file
    (animated GIF/WebP: first frame only).
    """
    from PIL import Image, ImageSequence

    _register_heif()
    multi = ext_of(path) in MULTIPAGE_IMAGE_EXTS
    paths: list[str] = []
    with Image.open(path) as im:
        frames = ImageSequence.Iterator(im) if multi else [im]
        for i, frame in enumerate(frames):
            if i >= max_pages:
                break
            out = os.path.join(out_dir, f"page_{i + 1:04d}.png")
            frame.convert("RGB").save(out, "PNG")
            paths.append(out)
    return paths


def image_preview_png(data: bytes, name: str) -> bytes:
    """First page/frame of an image as PNG bytes (so HEIC/TIFF/WebP can be shown in the browser)."""
    import io

    from PIL import Image

    _register_heif()
    with Image.open(io.BytesIO(data)) as im:
        buf = io.BytesIO()
        im.convert("RGB").save(buf, "PNG")
    return buf.getvalue()


def pdf_page_count(data: bytes) -> int:
    """Number of pages in a PDF given its bytes (PyMuPDF)."""
    import fitz  # PyMuPDF

    with fitz.open(stream=data, filetype="pdf") as doc:
        return doc.page_count


def pdf_page_png(data: bytes, page_index: int, dpi: int = 150) -> bytes:
    """Rasterise one page (0-based) of a PDF to PNG bytes with PyMuPDF.

    Used for the in-app preview: Chrome blocks data:/blob: PDF iframes on
    localhost, so we show a bitmap instead.
    """
    import fitz  # PyMuPDF

    with fitz.open(stream=data, filetype="pdf") as doc:
        if not 0 <= page_index < doc.page_count:
            raise IndexError(f"page {page_index + 1} out of range (1-{doc.page_count})")
        pix = doc[page_index].get_pixmap(dpi=dpi, alpha=False)
        return pix.tobytes("png")


_TESSERACT_DEFAULT_EXE = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def ocr_pages_tesseract(pages: list[str]) -> list[str]:
    import shutil
    import pytesseract
    from PIL import Image

    # The Windows installer does not add Tesseract to PATH; fall back to the
    # default install location so the free OCR mode works out of the box.
    if shutil.which(pytesseract.pytesseract.tesseract_cmd) is None and os.path.exists(_TESSERACT_DEFAULT_EXE):
        pytesseract.pytesseract.tesseract_cmd = _TESSERACT_DEFAULT_EXE

    texts = []
    for p in pages:
        with Image.open(p) as im:
            texts.append(pytesseract.image_to_string(im).strip())
    return texts


def join_pages(page_texts: list[str], name: str, header: str, byline: str | None = None) -> str:
    """Combine per-page Markdown into one document (page headings only when > 1 page)."""
    md = f"# {header}\n\n> Source file: {name}\n\n"
    if byline:
        md += f"> {byline}\n\n"
    if len(page_texts) == 1:
        return md + (page_texts[0] or "_No text found in this image._")
    for i, text in enumerate(page_texts, 1):
        md += f"## Page {i}\n\n{text or '_No text found on this page._'}\n\n"
        if i < len(page_texts):
            md += "---\n\n"
    return md.rstrip() + "\n"


# --------------------------------------------------------------------------- #
# AI enhancement (Claude / Ollama)
# --------------------------------------------------------------------------- #
def process_document_with_ai(
    raw_text: str,
    filename: str,
    doc_mode: str,
    anthropic_key: str = None,
    ollama_model_name: str = None
) -> str:
    """
    Enhances and structures raw Markdown text using Anthropic Claude AI or local Ollama.
    """
    import anthropic
    import ollama

    prompt = f"""You are an expert document analysis assistant.
We have converted a document named "{filename}" to raw Markdown text.
Your task is to structure, clean up, format, and enhance this Markdown.
- Keep the original information intact.
- Improve headings, formatting, tables, lists, and code blocks.
- Clean up any messy extraction artifacts.
- Provide ONLY the final, polished Markdown content. No conversational intro/outro.
Raw converted text:
---
{raw_text}
---
"""
    if doc_mode == DOC_MODE_CLAUDE:
        if not anthropic_key:
            raise ValueError("Anthropic API key is required but was not provided.")
        client = anthropic.Anthropic(api_key=anthropic_key)
        message = client.messages.create(
            model="claude-3-5-sonnet-20241022",
            max_tokens=4096,
            messages=[
                {"role": "user", "content": prompt}
            ]
        )
        return message.content[0].text
    elif doc_mode == DOC_MODE_OLLAMA:
        if not ollama_model_name:
            raise ValueError("Ollama model name is required but was not provided.")
        response = ollama.chat(
            model=ollama_model_name,
            messages=[
                {"role": "user", "content": prompt}
            ]
        )
        return strip_thinking(response['message']['content'])
    return raw_text


def strip_thinking(text: str) -> str:
    """Remove <think>...</think> reasoning blocks emitted by models such as deepseek-r1."""
    import re
    return re.sub(r"<think>.*?</think>\s*", "", text, flags=re.DOTALL).strip()


# --------------------------------------------------------------------------- #
# URL conversion (Wikipedia / YouTube / generic with fallbacks)
# --------------------------------------------------------------------------- #
class URLConversionError(Exception):
    """User-facing URL conversion failure."""


class YouTubeNoCaptionsError(URLConversionError):
    """YouTube video has no usable transcript — the UI offers a Whisper upload instead."""

    def __init__(self, message: str, detail: str = ""):
        super().__init__(message)
        self.detail = detail


# Caption language priority: English → Hindi → Marathi.
YOUTUBE_LANGUAGES = ("en", "hi", "mr")
YOUTUBE_NO_CAPTIONS_MSG = (
    "No captions available for this video. "
    "Upload the video file directly for Whisper transcription."
)
_YOUTUBE_ID_RE = re.compile(r"(?:v=|youtu\.be/|/shorts/|/embed/|/live/)([a-zA-Z0-9_-]{11})")


def extract_youtube_id(url: str) -> str | None:
    """Return the 11-character video id from a youtube.com / youtu.be URL, else None."""
    match = _YOUTUBE_ID_RE.search(url or "")
    return match.group(1) if match else None


def fetch_youtube_transcript(video_id: str, languages=YOUTUBE_LANGUAGES) -> tuple[str, dict]:
    """
    Fetch captions with youtube-transcript-api (>= 1.2) and return
    (markdown, sidecar). Markdown is one paragraph per caption cue:
    **[MM:SS]** text. Raises YouTubeNoCaptionsError when nothing usable exists.
    """
    from youtube_transcript_api import YouTubeTranscriptApi

    fetched = YouTubeTranscriptApi().fetch(video_id, languages=list(languages))
    segments = []
    for snip in fetched.snippets:
        text = " ".join((snip.text or "").split())
        if text:
            segments.append({"start": round(float(snip.start), 3), "duration": round(float(snip.duration), 3), "text": text})
    if not segments:
        raise YouTubeNoCaptionsError(YOUTUBE_NO_CAPTIONS_MSG, detail="transcript was empty")
    lines = [f"**[{int(seg['start'] // 60):02d}:{int(seg['start'] % 60):02d}]** {seg['text']}" for seg in segments]
    sidecar = {
        "video_id": video_id,
        "language": fetched.language,
        "language_code": fetched.language_code,
        "auto_generated": bool(fetched.is_generated),
        "segment_count": len(segments),
        "segments": segments,
    }
    return "\n\n".join(lines), sidecar


def fetch_url_content(url):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.5",
        "Accept-Encoding": "gzip, deflate, br",
        "Connection": "keep-alive",
    }
    response = requests.get(url, headers=headers, timeout=15)
    response.raise_for_status()
    return response.text


async def fetch_with_browser(url):
    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        page = await browser.new_page(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        )
        await page.goto(url, wait_until="networkidle", timeout=30000)
        html = await page.content()
        await browser.close()
        return html


def _friendly_url_error(exc: Exception) -> str:
    if isinstance(exc, requests.exceptions.ConnectionError):
        return "❌ Could not connect. Check your internet or the URL."
    if isinstance(exc, requests.exceptions.Timeout):
        return "❌ Request timed out. Website took too long to respond."
    text = str(exc)
    if "403" in text:
        return "❌ This website blocks automated access (403 Forbidden). Try a different URL."
    if "404" in text:
        return "❌ Page not found (404). Please check the URL."
    return f"❌ Conversion failed: {text}"


def convert_url(url_input: str, status=None) -> ConversionResult:
    """
    Convert a URL to Markdown. `status(message)` receives progress messages.
    Raises URLConversionError with a user-facing message on failure.
    """
    from markitdown import MarkItDown

    say = status or (lambda _msg: None)
    try:
        if "wikipedia.org/wiki/" in url_input:
            import wikipediaapi

            page_title = url_input.split("/wiki/")[-1].replace("_", " ")
            lang = url_input.split("//")[-1].split(".")[0]
            wiki = wikipediaapi.Wikipedia(user_agent="MarkitDownApp/1.0", language=lang)
            page = wiki.page(page_title)
            if not page.exists():
                raise URLConversionError("❌ Wikipedia page not found. Check the URL.")
            markdown_text = f"# {page.title}\n\n> source: {url_input}\n\n{page.text}"
            return ConversionResult(markdown_text, page_title, "url", sidecar={"url": url_input, "site": "wikipedia"})

        # ---- YouTube: explicit caption fetch, checked BEFORE the MarkItDown default ----
        video_id = extract_youtube_id(url_input)
        if video_id or "youtube.com/" in url_input or "youtu.be/" in url_input:
            if not video_id:
                raise URLConversionError("❌ Could not find a video id in that YouTube URL.")
            say("⏳ Fetching YouTube captions (English → Hindi → Marathi)...")
            try:
                body, sidecar = fetch_youtube_transcript(video_id)
            except URLConversionError:
                raise
            except Exception as exc:
                # TranscriptsDisabled / NoTranscriptFound / VideoUnavailable / network — all end here.
                raise YouTubeNoCaptionsError(YOUTUBE_NO_CAPTIONS_MSG, detail=f"{type(exc).__name__}: {exc}") from exc
            markdown_text = f"# YouTube transcript — {video_id}\n\n> source: {url_input}\n\n{body}"
            sidecar.update({"url": url_input, "site": "youtube"})
            return ConversionResult(markdown_text, f"youtube_{video_id}", "url", sidecar=sidecar)

        tmp_html_path = None
        try:
            md = MarkItDown()
            # ---- Attempt 1: MarkItDown direct ----
            raw_text = md.convert(url_input).text_content
            # ---- Attempt 2: requests with browser headers ----
            if not raw_text.strip():
                say("⏳ Trying with browser headers...")
                html_content = fetch_url_content(url_input)
                with tempfile.NamedTemporaryFile(delete=False, suffix=".html", mode="w", encoding="utf-8") as tmp_html:
                    tmp_html.write(html_content)
                    tmp_html_path = tmp_html.name
                raw_text = md.convert(tmp_html_path).text_content
            # ---- Attempt 3: Playwright headless browser ----
            if not raw_text.strip():
                say("⏳ Trying headless browser for JavaScript rendered page...")
                html_content = asyncio.run(fetch_with_browser(url_input))
                with tempfile.NamedTemporaryFile(delete=False, suffix=".html", mode="w", encoding="utf-8") as tmp_html:
                    tmp_html.write(html_content)
                    tmp_html_path = tmp_html.name
                raw_text = md.convert(tmp_html_path).text_content
            if not raw_text.strip():
                raise URLConversionError("❌ Could not extract content. Page may require login or is fully blocked.")
            return ConversionResult(raw_text, "webpage", "url", sidecar={"url": url_input, "site": "generic"})
        finally:
            if tmp_html_path and os.path.exists(tmp_html_path):
                os.remove(tmp_html_path)
    except URLConversionError:
        raise
    except Exception as exc:
        raise URLConversionError(_friendly_url_error(exc)) from exc
